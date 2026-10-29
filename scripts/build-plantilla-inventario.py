"""Plantilla gratuita de control de inventario (Neurovia Systems).

Se descarga desde /guias/plantilla-de-inventario-en-excel.

Hojas: Cómo usar · Resumen · Productos · Movimientos · Kardex.
Solo funciones de Excel 2007 (SUMIFS, INDEX/MATCH, COUNTIF, EDATE, IFERROR) para
que abra igual en Excel 2010+, Excel 365, Google Sheets y LibreOffice.

Uso:
    python scripts/build-plantilla-inventario.py [salida.xlsx]

openpyxl guarda las fórmulas sin resultados. Antes de publicar, ábrelo en Excel
para que calcule y guarde los valores (si no, las vistas previas salen vacías):

    $xl = New-Object -ComObject Excel.Application; $xl.DisplayAlerts = $false
    $wb = $xl.Workbooks.Open("<ruta absoluta>"); $xl.CalculateFull()
    $wb.Worksheets.Item(1).Activate(); $wb.Save(); $wb.Close($false); $xl.Quit()
"""
import sys
from datetime import date

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OUT = sys.argv[1] if len(sys.argv) > 1 else "public/descargas/plantilla-control-de-inventario.xlsx"

FIRST = 5                 # primera fila de datos en Productos y Movimientos
P_LAST = FIRST + 500 - 1  # 500 artículos
M_LAST = FIRST + 3000 - 1 # 3,000 movimientos
K_ROWS = 300              # renglones del kardex
R_ROWS = 25               # renglones de la lista de compras

UTM = "?utm_source=plantilla-inventario&utm_medium=excel&utm_campaign=plantillas"
SITE = "https://neuroviasystems.com.mx"
GUIA = f"{SITE}/guias/plantilla-de-inventario-en-excel{UTM}"
SISTEMA = f"{SITE}/productos/software-de-inventario{UTM}"

FONT = "Arial"
ACCENT = "2F8FE8"
INK = "1E293B"
INPUT_FILL = PatternFill("solid", fgColor="FFF8E1")   # se captura
CALC_FILL = PatternFill("solid", fgColor="F1F5F9")    # fórmula, no tocar
HEAD_FILL = PatternFill("solid", fgColor=INK)
thin = Side(style="thin", color="D9DEE5")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

MONEY = '"$"#,##0.00'
DATE = "dd/mm/yyyy"


def f(size=10, bold=False, color=INK, underline=None, italic=False):
    return Font(name=FONT, size=size, bold=bold, color=color, underline=underline, italic=italic)


def title(ws, text, subtitle):
    ws["A1"] = text
    ws["A1"].font = f(16, True)
    ws["A2"] = subtitle
    ws["A2"].font = f(9, color="64748B")
    ws.row_dimensions[1].height = 24


def header(ws, row, cols):
    for i, (name, width, kind) in enumerate(cols, start=1):
        c = ws.cell(row=row, column=i, value=name)
        c.font = f(10, True, "FFFFFF")
        c.fill = HEAD_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[row].height = 30


def style_body(ws, first, last, cols):
    for i, (_, _, kind) in enumerate(cols, start=1):
        fill = INPUT_FILL if kind == "in" else CALC_FILL
        for r in range(first, last + 1):
            c = ws.cell(row=r, column=i)
            c.fill = fill
            c.border = BORDER
            c.font = f(10)


wb = Workbook()

# ---------------------------------------------------------------- Productos
ps = wb.active
ps.title = "Productos"
P_COLS = [
    ("Clave", 11, "in"),
    ("Descripción", 34, "in"),
    ("Categoría", 15, "in"),
    ("Unidad", 12, "in"),
    ("Ubicación", 13, "in"),
    ("Costo unitario", 13, "in"),
    ("Stock mínimo", 10, "in"),
    ("Stock máximo", 10, "in"),
    ("Existencia inicial", 11, "in"),
    ("Entradas", 10, "fx"),
    ("Salidas", 10, "fx"),
    ("Existencia actual", 11, "fx"),
    ("Valor del inventario", 15, "fx"),
    ("Estado", 12, "fx"),
    ("n_compra", 6, "fx"),  # auxiliar para la lista de compras (oculta)
]
title(ps, "Catálogo y existencias", "Captura tus artículos en las columnas amarillas. Las grises se calculan solas con los movimientos.")
header(ps, 4, P_COLS)
style_body(ps, FIRST, P_LAST, P_COLS)

MV = f"Movimientos!$F${FIRST}:$F${M_LAST}"
MC = f"Movimientos!$C${FIRST}:$C${M_LAST}"
MB = f"Movimientos!$B${FIRST}:$B${M_LAST}"
for r in range(FIRST, P_LAST + 1):
    ps[f"J{r}"] = f'=IF($A{r}="","",SUMIFS({MV},{MC},$A{r},{MB},"Entrada"))'
    ps[f"K{r}"] = f'=IF($A{r}="","",SUMIFS({MV},{MC},$A{r},{MB},"Salida"))'
    ps[f"L{r}"] = f'=IF($A{r}="","",N(I{r})+J{r}-K{r})'
    ps[f"M{r}"] = f'=IF($A{r}="","",L{r}*N(F{r}))'
    ps[f"N{r}"] = f'=IF($A{r}="","",IF(L{r}<=0,"Agotado",IF(AND(N(G{r})>0,L{r}<=G{r}),"Reordenar","OK")))'
    ps[f"O{r}"] = f'=IF(OR(N{r}="Reordenar",N{r}="Agotado"),COUNTIF($N${FIRST}:N{r},"Reordenar")+COUNTIF($N${FIRST}:N{r},"Agotado"),"")'
    ps[f"F{r}"].number_format = MONEY
    ps[f"M{r}"].number_format = MONEY
    ps[f"N{r}"].alignment = Alignment(horizontal="center")
ps.column_dimensions["O"].hidden = True
ps.freeze_panes = f"C{FIRST}"
ps.auto_filter.ref = f"A4:N{P_LAST}"

dv_key = DataValidation(type="custom", formula1=f"COUNTIF($A${FIRST}:$A${P_LAST},A{FIRST})=1", allow_blank=True,
                        showErrorMessage=True, errorTitle="Clave repetida",
                        error="Esa clave ya existe en el catálogo. Cada artículo necesita una clave única.")
dv_key.add(f"A{FIRST}:A{P_LAST}")
dv_num = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True,
                        showErrorMessage=True, errorTitle="Número inválido", error="Escribe un número igual o mayor a cero.")
dv_num.add(f"F{FIRST}:I{P_LAST}")
ps.add_data_validation(dv_key)
ps.add_data_validation(dv_num)

rng = f"N{FIRST}:N{P_LAST}"
ps.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Agotado"'], fill=PatternFill("solid", fgColor="FDE2E1"), font=Font(name=FONT, bold=True, color="B42318")))
ps.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Reordenar"'], fill=PatternFill("solid", fgColor="FEF0C7"), font=Font(name=FONT, bold=True, color="B54708")))
ps.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"OK"'], font=Font(name=FONT, color="067647")))

productos = [
    ("P-001", "Cemento gris 50 kg", "Materiales", "Bulto", "Bodega A", 215, 40, 120, 80),
    ("P-002", "Varilla corrugada 3/8\"", "Acero", "Pieza", "Patio", 168, 100, 300, 150),
    ("P-003", "Block de concreto 15x20x40", "Materiales", "Pieza", "Patio", 14.5, 500, 2000, 1200),
    ("P-004", "Tubo PVC hidráulico 1/2\"", "Plomería", "Tramo", "Bodega B", 62, 30, 100, 45),
    ("P-005", "Pintura vinílica blanca 19 L", "Pinturas", "Cubeta", "Bodega B", 1150, 5, 20, 8),
    ("P-006", "Alambre recocido", "Acero", "Kg", "Bodega A", 38, 50, 200, 60),
    ("P-007", "Cable THW calibre 12", "Eléctrico", "Rollo 100 m", "Bodega B", 980, 4, 12, 3),
    ("P-008", "Tornillo para tablaroca 1\"", "Ferretería", "Caja 1000", "Bodega A", 145, 10, 40, 22),
]
for i, row in enumerate(productos):
    for j, v in enumerate(row):
        ps.cell(row=FIRST + i, column=j + 1, value=v)
ps[f"A{FIRST}"].comment = Comment("Datos de ejemplo: bórralos (columnas A a I) cuando captures tu catálogo.", "Neurovia Systems")

# ---------------------------------------------------------------- Movimientos
ms = wb.create_sheet("Movimientos")
M_COLS = [
    ("Fecha", 12, "in"),
    ("Tipo", 10, "in"),
    ("Clave", 11, "in"),
    ("Descripción", 32, "fx"),
    ("Unidad", 11, "fx"),
    ("Cantidad", 10, "in"),
    ("Costo unitario", 13, "fx"),
    ("Importe", 14, "fx"),
    ("Existencia después", 12, "fx"),
    ("Referencia", 16, "in"),
    ("Responsable", 15, "in"),
    ("Notas", 28, "in"),
    ("n_kardex", 6, "fx"),  # auxiliar para el kardex (oculta)
]
title(ms, "Entradas y salidas", "Una fila por movimiento, en orden de fecha. Tipo: Entrada o Salida. La clave debe existir en Productos.")
header(ms, 4, M_COLS)
style_body(ms, FIRST, M_LAST, M_COLS)

PA = f"Productos!$A${FIRST}:$A${P_LAST}"
for r in range(FIRST, M_LAST + 1):
    match = f"MATCH($C{r},{PA},0)"
    ms[f"D{r}"] = f'=IF($C{r}="","",IFERROR(INDEX(Productos!$B${FIRST}:$B${P_LAST},{match}),"Clave no existe"))'
    ms[f"E{r}"] = f'=IF($C{r}="","",IFERROR(INDEX(Productos!$D${FIRST}:$D${P_LAST},{match}),""))'
    ms[f"G{r}"] = f'=IF($C{r}="","",IFERROR(INDEX(Productos!$F${FIRST}:$F${P_LAST},{match}),0))'
    ms[f"H{r}"] = f'=IF(OR($C{r}="",$F{r}=""),"",F{r}*G{r})'
    ms[f"I{r}"] = (
        f'=IF(OR($C{r}="",$F{r}=""),"",IFERROR(INDEX(Productos!$I${FIRST}:$I${P_LAST},{match}),0)'
        f'+SUMIFS($F${FIRST}:$F{r},$C${FIRST}:$C{r},$C{r},$B${FIRST}:$B{r},"Entrada")'
        f'-SUMIFS($F${FIRST}:$F{r},$C${FIRST}:$C{r},$C{r},$B${FIRST}:$B{r},"Salida"))'
    )
    ms[f"M{r}"] = f'=IF(AND($C{r}<>"",$C{r}=Kardex!$C$4),COUNTIF($C${FIRST}:$C{r},Kardex!$C$4),"")'
    ms[f"A{r}"].number_format = DATE
    ms[f"G{r}"].number_format = MONEY
    ms[f"H{r}"].number_format = MONEY
ms.column_dimensions["M"].hidden = True
ms.freeze_panes = f"D{FIRST}"
ms.auto_filter.ref = f"A4:L{M_LAST}"

dv_tipo = DataValidation(type="list", formula1='"Entrada,Salida"', allow_blank=True, showErrorMessage=True,
                         errorTitle="Tipo inválido", error="Elige Entrada o Salida.")
dv_tipo.add(f"B{FIRST}:B{M_LAST}")
dv_clave = DataValidation(type="list", formula1=f"={PA}", allow_blank=True, showErrorMessage=True,
                          errorTitle="Clave no registrada", error="Primero da de alta el artículo en la hoja Productos.")
dv_clave.add(f"C{FIRST}:C{M_LAST}")
dv_cant = DataValidation(type="decimal", operator="greaterThan", formula1="0", allow_blank=True, showErrorMessage=True,
                         errorTitle="Cantidad inválida", error="La cantidad va en positivo; el tipo (Entrada/Salida) define si suma o resta.")
dv_cant.add(f"F{FIRST}:F{M_LAST}")
dv_fecha = DataValidation(type="date", operator="greaterThan", formula1="36526", allow_blank=True, showErrorMessage=True,
                          errorTitle="Fecha inválida", error="Escribe una fecha, por ejemplo 15/10/2026.")
dv_fecha.add(f"A{FIRST}:A{M_LAST}")
for dv in (dv_tipo, dv_clave, dv_cant, dv_fecha):
    ms.add_data_validation(dv)

ms.conditional_formatting.add(
    f"I{FIRST}:I{M_LAST}",
    FormulaRule(formula=[f'AND(ISNUMBER(I{FIRST}),I{FIRST}<0)'], fill=PatternFill("solid", fgColor="FDE2E1"), font=Font(name=FONT, bold=True, color="B42318")),
)
ms.conditional_formatting.add(
    f"D{FIRST}:D{M_LAST}",
    CellIsRule(operator="equal", formula=['"Clave no existe"'], font=Font(name=FONT, bold=True, color="B42318")),
)

movs = [
    (date(2026, 10, 1), "Entrada", "P-001", 60, "Factura A-1182", "Almacén"),
    (date(2026, 10, 2), "Salida", "P-001", 35, "Pedido 0451", "Ventas"),
    (date(2026, 10, 3), "Salida", "P-003", 400, "Pedido 0452", "Ventas"),
    (date(2026, 10, 6), "Entrada", "P-002", 120, "Factura B-0907", "Almacén"),
    (date(2026, 10, 7), "Salida", "P-002", 210, "Pedido 0458", "Ventas"),
    (date(2026, 10, 9), "Salida", "P-005", 4, "Pedido 0461", "Ventas"),
    (date(2026, 10, 10), "Salida", "P-001", 50, "Pedido 0463", "Ventas"),
    (date(2026, 10, 13), "Entrada", "P-004", 40, "Factura C-2210", "Almacén"),
    (date(2026, 10, 14), "Salida", "P-004", 52, "Pedido 0470", "Ventas"),
    (date(2026, 10, 16), "Salida", "P-007", 3, "Pedido 0474", "Ventas"),
    (date(2026, 10, 20), "Salida", "P-006", 25, "Pedido 0480", "Ventas"),
    (date(2026, 10, 21), "Entrada", "P-008", 10, "Factura A-1240", "Almacén"),
    (date(2026, 10, 23), "Salida", "P-008", 14, "Pedido 0486", "Ventas"),
    (date(2026, 10, 27), "Salida", "P-003", 380, "Pedido 0491", "Ventas"),
]
for i, (d, tipo, clave, cant, ref, resp) in enumerate(movs):
    r = FIRST + i
    ms[f"A{r}"] = d
    ms[f"B{r}"] = tipo
    ms[f"C{r}"] = clave
    ms[f"F{r}"] = cant
    ms[f"J{r}"] = ref
    ms[f"K{r}"] = resp
ms[f"A{FIRST}"].comment = Comment("Datos de ejemplo: bórralos (columnas A, B, C, F, J, K y L) antes de capturar los tuyos.", "Neurovia Systems")

# ---------------------------------------------------------------- Kardex
ks = wb.create_sheet("Kardex")
title(ks, "Kardex por artículo", "Elige una clave en la celda amarilla y verás todos sus movimientos con la existencia después de cada uno.")
ks["B4"] = "Clave:"
ks["B4"].font = f(11, True)
ks["B4"].alignment = Alignment(horizontal="right")
ks["C4"] = "P-001"
ks["C4"].fill = INPUT_FILL
ks["C4"].font = f(11, True)
ks["C4"].border = BORDER
ks["D4"] = f'=IF($C$4="","",IFERROR(INDEX(Productos!$B${FIRST}:$B${P_LAST},MATCH($C$4,{PA},0)),"Clave no existe"))'
ks["D4"].font = f(11, True)
dv_k = DataValidation(type="list", formula1=f"={PA}", allow_blank=True)
dv_k.add("C4")
ks.add_data_validation(dv_k)

ks["B5"] = "Existencia inicial:"
ks["C5"] = f'=IF($C$4="","",IFERROR(INDEX(Productos!$I${FIRST}:$I${P_LAST},MATCH($C$4,{PA},0)),0))'
ks["B6"] = "Existencia actual:"
ks["C6"] = f'=IF($C$4="","",IFERROR(INDEX(Productos!$L${FIRST}:$L${P_LAST},MATCH($C$4,{PA},0)),0))'
for a in ("B5", "B6"):
    ks[a].font = f(10, True)
    ks[a].alignment = Alignment(horizontal="right")
for a in ("C5", "C6"):
    ks[a].font = f(10, True)
    ks[a].fill = CALC_FILL
    ks[a].border = BORDER

K_COLS = [("#", 5, "fx"), ("Fecha", 12, "fx"), ("Tipo", 10, "fx"), ("Referencia", 18, "fx"), ("Entrada", 10, "fx"),
          ("Salida", 10, "fx"), ("Existencia", 11, "fx"), ("Costo unitario", 13, "fx"), ("Importe", 14, "fx"), ("Responsable", 15, "fx")]
KH = 8
header(ks, KH, K_COLS)
style_body(ks, KH + 1, KH + K_ROWS, K_COLS)
MK = f"Movimientos!$M${FIRST}:$M${M_LAST}"


def mcol(col, n):
    return f"INDEX(Movimientos!${col}${FIRST}:${col}${M_LAST},MATCH({n},{MK},0))"


for i in range(1, K_ROWS + 1):
    r = KH + i
    n = f"$A{r}"
    ks[f"A{r}"] = f'=IF({i}<=COUNTIF({MK},">0"),{i},"")'
    ks[f"B{r}"] = f'=IF({n}="","",{mcol("A", n)})'
    ks[f"C{r}"] = f'=IF({n}="","",{mcol("B", n)})'
    ks[f"D{r}"] = f'=IF({n}="","",{mcol("J", n)}&"")'
    ks[f"E{r}"] = f'=IF({n}="","",IF(C{r}="Entrada",{mcol("F", n)},""))'
    ks[f"F{r}"] = f'=IF({n}="","",IF(C{r}="Salida",{mcol("F", n)},""))'
    ks[f"G{r}"] = f'=IF({n}="","",{mcol("I", n)})'
    ks[f"H{r}"] = f'=IF({n}="","",{mcol("G", n)})'
    ks[f"I{r}"] = f'=IF({n}="","",{mcol("H", n)})'
    ks[f"J{r}"] = f'=IF({n}="","",{mcol("K", n)}&"")'
    ks[f"B{r}"].number_format = DATE
    ks[f"H{r}"].number_format = MONEY
    ks[f"I{r}"].number_format = MONEY
ks.freeze_panes = f"A{KH + 1}"

# ---------------------------------------------------------------- Resumen
rs = wb.create_sheet("Resumen", 0)
title(rs, "Resumen del inventario", "Se actualiza solo con lo que captures en Productos y Movimientos.")
rs.column_dimensions["A"].width = 3
for col, w in zip("BCDEFGH", (30, 18, 34, 12, 12, 12, 14)):
    rs.column_dimensions[col].width = w

rs["B4"] = "Mes a revisar"
rs["C4"] = date(2026, 10, 1)
rs["C4"].number_format = "mmmm yyyy"
rs["C4"].fill = INPUT_FILL
rs["C4"].border = BORDER
rs["C4"].comment = Comment("Escribe cualquier fecha del mes que quieres revisar (por ejemplo 01/11/2026).", "Neurovia Systems")
rs["C4"].font = f(10, True)
rs["B4"].font = f(10, True)

PN = f"Productos!$N${FIRST}:$N${P_LAST}"
ini = 'DATE(YEAR($C$4),MONTH($C$4),1)'
MA = f"Movimientos!$A${FIRST}:$A${M_LAST}"
MH = f"Movimientos!$H${FIRST}:$H${M_LAST}"
kpis = [
    ("Artículos en el catálogo", f"=COUNTA({PA})", "0"),
    ("Valor total del inventario", f"=SUM(Productos!$M${FIRST}:$M${P_LAST})", MONEY),
    ("Artículos por reordenar", f'=COUNTIF({PN},"Reordenar")', "0"),
    ("Artículos agotados", f'=COUNTIF({PN},"Agotado")', "0"),
    ("Movimientos de entrada en el mes", f'=COUNTIFS({MB},"Entrada",{MA},">="&{ini},{MA},"<"&EDATE({ini},1))', "0"),
    ("Movimientos de salida en el mes", f'=COUNTIFS({MB},"Salida",{MA},">="&{ini},{MA},"<"&EDATE({ini},1))', "0"),
    ("Valor de las entradas del mes", f'=SUMIFS({MH},{MB},"Entrada",{MA},">="&{ini},{MA},"<"&EDATE({ini},1))', MONEY),
    ("Valor de las salidas del mes", f'=SUMIFS({MH},{MB},"Salida",{MA},">="&{ini},{MA},"<"&EDATE({ini},1))', MONEY),
]
for i, (label, formula, fmt) in enumerate(kpis):
    r = 6 + i
    rs[f"B{r}"] = label
    rs[f"B{r}"].font = f(10)
    rs[f"C{r}"] = formula
    rs[f"C{r}"].number_format = fmt
    rs[f"C{r}"].font = f(12, True)
    rs[f"C{r}"].alignment = Alignment(horizontal="right")
    for col in "BC":
        rs[f"{col}{r}"].border = BORDER
        rs[f"{col}{r}"].fill = CALC_FILL
rs.conditional_formatting.add("C8:C9", CellIsRule(operator="greaterThan", formula=["0"], font=Font(name=FONT, size=12, bold=True, color="B42318")))

LR = 16
rs[f"B{LR - 1}"] = "Lista de compras sugerida"
rs[f"B{LR - 1}"].font = f(13, True)
rs[f"D{LR - 1}"] = "Artículos agotados o en su mínimo. Sugerido = llevarlo al stock máximo (o al doble del mínimo si no capturaste máximo)."
rs[f"D{LR - 1}"].font = f(8, italic=True, color="64748B")
R_COLS = [("Clave", "C"), ("Descripción", "D"), ("Existencia", "E"), ("Mínimo", "F"), ("Estado", "G"), ("Sugerido comprar", "H")]
rs[f"B{LR}"] = "#"
for name, col in [("#", "B")] + R_COLS:
    c = rs[f"{col}{LR}"]
    c.value = name
    c.font = f(10, True, "FFFFFF")
    c.fill = HEAD_FILL
    c.border = BORDER
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
PO = f"Productos!$O${FIRST}:$O${P_LAST}"


def pcol(col, n):
    return f"INDEX(Productos!${col}${FIRST}:${col}${P_LAST},MATCH({n},{PO},0))"


for i in range(1, R_ROWS + 1):
    r = LR + i
    n = f"$B{r}"
    rs[f"B{r}"] = f'=IF({i}<=COUNTIF({PO},">0"),{i},"")'
    rs[f"C{r}"] = f'=IF({n}="","",{pcol("A", n)})'
    rs[f"D{r}"] = f'=IF({n}="","",{pcol("B", n)})'
    rs[f"E{r}"] = f'=IF({n}="","",{pcol("L", n)})'
    rs[f"F{r}"] = f'=IF({n}="","",{pcol("G", n)})'
    rs[f"G{r}"] = f'=IF({n}="","",{pcol("N", n)})'
    rs[f"H{r}"] = f'=IF({n}="","",MAX(0,IF(N({pcol("H", n)})>0,N({pcol("H", n)}),2*N(F{r}))-E{r}))'
    for col in "BCDEFGH":
        rs[f"{col}{r}"].border = BORDER
        rs[f"{col}{r}"].fill = CALC_FILL
        rs[f"{col}{r}"].font = f(10)
rs[f"B{LR + R_ROWS + 1}"] = f"Muestra hasta {R_ROWS} artículos. Si tienes más pendientes, filtra la columna Estado en la hoja Productos."
rs[f"B{LR + R_ROWS + 1}"].font = f(8, italic=True, color="64748B")
rs.conditional_formatting.add(f"G{LR + 1}:G{LR + R_ROWS}", CellIsRule(operator="equal", formula=['"Agotado"'], font=Font(name=FONT, bold=True, color="B42318")))
rs.conditional_formatting.add(f"G{LR + 1}:G{LR + R_ROWS}", CellIsRule(operator="equal", formula=['"Reordenar"'], font=Font(name=FONT, bold=True, color="B54708")))

# ---------------------------------------------------------------- Cómo usar
hs = wb.create_sheet("Cómo usar", 0)
hs.column_dimensions["A"].width = 3
hs.column_dimensions["B"].width = 100
hs["B2"] = "Plantilla de control de inventario"
hs["B2"].font = f(20, True)
hs["B3"] = "Entradas, salidas, existencias, alertas de reorden y kardex por artículo. Gratis, sin macros."
hs["B3"].font = f(11, color="64748B")

lines = [
    ("Cómo se usa", "h"),
    ("1. En Productos, da de alta cada artículo con una clave única, su costo, su stock mínimo (y máximo, si lo manejas) y la existencia con la que arrancas.", "p"),
    ("2. En Movimientos, captura una fila por cada entrada (compra, devolución de cliente) o salida (venta, consumo, merma), en orden de fecha.", "p"),
    ("3. Las existencias, el valor del inventario y el estado (OK, Reordenar, Agotado) se calculan solos en Productos.", "p"),
    ("4. En Resumen ves los totales del mes y la lista de compras sugerida; en Kardex, el historial completo de un artículo.", "p"),
    ("5. Los datos que trae son de ejemplo: bórralos de Productos (columnas A a I) y de Movimientos (A, B, C, F, J, K y L) antes de empezar.", "p"),
    ("Colores", "h"),
    ("Celdas amarillas: aquí capturas.", "in"),
    ("Celdas grises: fórmulas. No las borres ni escribas encima; si se rompe una, cópiala de la fila de arriba.", "fx"),
    ("Buenas prácticas", "h"),
    ("• Las cantidades van siempre en positivo: el tipo (Entrada o Salida) decide si suma o resta.", "p"),
    ("• Haz un conteo físico al menos una vez al mes y registra la diferencia como Entrada o Salida con la referencia \"Ajuste por conteo\".", "p"),
    ("• Si la Existencia después sale en rojo, se registró una salida mayor a lo que había: revisa ese movimiento.", "p"),
    ("• Capacidad: 500 artículos y 3,000 movimientos. Para más, copia la última fila con fórmulas hacia abajo.", "p"),
    ("Cuándo Excel ya no alcanza", "h"),
    ("Si capturan varias personas a la vez, manejas más de un almacén, necesitas lotes y caducidades o quieres que la venta descuente sola el inventario, una hoja de cálculo empieza a fallar. Ahí conviene un sistema.", "p"),
    ("Guía completa: cómo llevar el inventario en Excel y cuándo cambiar a un sistema", GUIA),
    ("Sistema de inventario de Neurovia Systems", SISTEMA),
    ("", "p"),
    ("Plantilla gratuita de Neurovia Systems · Villahermosa, Tabasco · neuroviasystems.com.mx · WhatsApp 993 722 6350", "foot"),
    ("Puedes usarla y compartirla libremente; si la publicas en tu sitio, enlaza a la guía de arriba.", "foot"),
]
r = 5
for text, kind in lines:
    c = hs[f"B{r}"]
    c.value = text
    c.alignment = Alignment(wrap_text=True, vertical="top")
    if kind == "h":
        r += 1
        c = hs[f"B{r}"]
        hs[f"B{r - 1}"].value = None
        c.value = text
        c.font = f(13, True, ACCENT)
    elif kind == "in":
        c.font = f(10)
        c.fill = INPUT_FILL
        c.border = BORDER
    elif kind == "fx":
        c.font = f(10)
        c.fill = CALC_FILL
        c.border = BORDER
    elif kind == "foot":
        c.font = f(9, color="64748B")
    elif kind.startswith("http"):
        c.hyperlink = kind
        c.font = f(10, color=ACCENT, underline="single")
    else:
        c.font = f(10)
    if kind != "h" and text:
        # Excel no ajusta solo el alto de las filas con texto ajustado: ~85 caracteres por renglón.
        hs.row_dimensions[r].height = 14.5 * max(1, -(-len(text) // 85))
    r += 1
hs.sheet_view.showGridLines = False
rs.sheet_view.showGridLines = False
ks.sheet_view.showGridLines = False

for ws in wb.worksheets:
    ws.sheet_properties.tabColor = {"Cómo usar": ACCENT, "Resumen": "067647"}.get(ws.title, "94A3B8")

wb.properties.title = "Plantilla de control de inventario"
wb.properties.creator = "Neurovia Systems"
wb.properties.subject = "Control de inventario en Excel"
wb.properties.keywords = "inventario, excel, kardex, control de almacén"
wb.active = 0
wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print("ok", OUT)
